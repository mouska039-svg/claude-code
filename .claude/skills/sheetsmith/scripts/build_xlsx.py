#!/usr/bin/env python3
"""Rend une spécification de classeur Sheetsmith en .xlsx (openpyxl).

Usage: python build_xlsx.py <spec.json> <sortie.xlsx>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import column_index_from_string

from spec_common import load_spec, col_letter, resolve_formula, data_bounds

THEME = {
    "primary": "5C7A6B", "accent": "D4876B", "bg": "F5EFE6",
    "muted": "8B8A87", "ink": "2C2C2A", "alt": "F5EFE6",
}


def hexn(c):
    return c.lstrip("#").upper()


def build_data_tab(ws, tab):
    cols = tab.get("columns", [])
    firstrow, lastrow = data_bounds(tab)
    header_fill = PatternFill("solid", fgColor=THEME["primary"])
    header_font = Font(bold=True, color="FFFFFF")
    for ci, col in enumerate(cols):
        cell = ws.cell(row=1, column=ci + 1, value=col.get("header", ""))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        if col.get("width"):
            ws.column_dimensions[col_letter(ci)].width = col["width"] / 7.0

    rows = tab.get("rows", []) or []
    for ri, rowdata in enumerate(rows):
        excel_row = firstrow + ri
        for ci, col in enumerate(cols):
            key = col.get("key")
            if col.get("formula"):
                val = resolve_formula(col["formula"], row=excel_row,
                                      firstrow=firstrow, lastrow=lastrow)
            else:
                val = rowdata.get(key)
            cell = ws.cell(row=excel_row, column=ci + 1, value=val)
            if col.get("format"):
                cell.number_format = col["format"]
            if ri % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=THEME["alt"])

    tr = tab.get("totals_row")
    if tr:
        trow = lastrow + 1
        ws.cell(row=trow, column=1, value=tr.get("label", "Total")).font = Font(bold=True)
        for ci, col in enumerate(cols):
            f = tr.get("formulas", {}).get(col.get("key"))
            if f:
                cell = ws.cell(row=trow, column=ci + 1,
                               value=resolve_formula(f, firstrow=firstrow, lastrow=lastrow))
                cell.font = Font(bold=True)
                if col.get("format"):
                    cell.number_format = col["format"]

    fr = tab.get("freeze")
    if fr:
        ws.freeze_panes = ws.cell(row=fr.get("rows", 0) + 1, column=fr.get("cols", 0) + 1)

    for cf in tab.get("conditional_formats", []):
        rng = resolve_formula(cf["range"], firstrow=firstrow, lastrow=lastrow)
        if cf["type"] == "color_scale":
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color=hexn(cf.get("min", "D4876B")),
                mid_type="percentile", mid_value=50, mid_color=hexn(cf.get("mid", "F5EFE6")),
                end_type="max", end_color=hexn(cf.get("max", "5C7A6B"))))
        elif cf["type"] == "data_bar":
            ws.conditional_formatting.add(rng, DataBarRule(
                start_type="min", end_type="max", color=hexn(cf.get("color", "5C7A6B"))))


def add_charts(ws, tab, wb):
    for ch in tab.get("charts", []):
        firstrow, lastrow = 1, 50
        chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(
            ch["type"], BarChart)()
        chart.title = ch.get("title", "")
        # data_range "Onglet!A1:C13"
        rng = ch["data_range"]
        sheet_name, a1 = rng.split("!")
        src = wb[sheet_name]
        start, end = a1.split(":")
        c1 = column_index_from_string("".join(filter(str.isalpha, start)))
        r1 = int("".join(filter(str.isdigit, start)))
        c2 = column_index_from_string("".join(filter(str.isalpha, end)))
        r2 = int("".join(filter(str.isdigit, end)))
        data = Reference(src, min_col=c1 + 1, max_col=c2, min_row=r1, max_row=r2)
        cats = Reference(src, min_col=c1, max_col=c1, min_row=r1 + 1, max_row=r2)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, ch.get("anchor", "H2"))


def build_dashboard(ws, tab, wb):
    widths = tab.get("column_widths", [])
    for ci, w in enumerate(widths):
        ws.column_dimensions[col_letter(ci)].width = w / 7.0
    for bar in tab.get("section_bars", []):
        rng = bar["range"]
        fill = PatternFill("solid", fgColor=hexn(bar.get("fill", THEME["primary"])))
        start = rng.split(":")[0]
        first_cell = ws[start]
        first_cell.value = bar.get("label", "")
        first_cell.font = Font(bold=True, color=hexn(bar.get("text", "FFFFFF")))
        for row in ws[rng]:
            for c in row:
                c.fill = fill
        ws.merge_cells(rng)
    for p in tab.get("panels", []):
        fill = PatternFill("solid", fgColor=hexn(p.get("fill", "FFFFFF")))
        for row in ws[p["range"]]:
            for c in row:
                c.fill = fill
    for card in tab.get("kpi_cards", []):
        anchor = card["anchor"]
        ac = ws[anchor]
        ac.value = card.get("label", "")
        ac.font = Font(size=9, color=hexn(THEME["muted"]))
        # grand nombre une ligne en dessous
        r = ac.row + 1
        cidx = ac.column
        vc = ws.cell(row=r, column=cidx, value=card.get("value"))
        vc.font = Font(size=20, bold=True, color=hexn(card.get("accent", THEME["primary"])))
        if card.get("format"):
            vc.number_format = card["format"]
    add_charts(ws, tab, wb)


def main():
    if len(sys.argv) < 3:
        raise SystemExit("Usage: build_xlsx.py <spec.json> <sortie.xlsx>")
    spec = load_spec(sys.argv[1])
    out = sys.argv[2]
    wb = Workbook()
    wb.remove(wb.active)
    # créer d'abord tous les onglets pour que les références croisées résolvent
    for tab in spec["tabs"]:
        wb.create_sheet(title=tab["name"])
    for tab in spec["tabs"]:
        ws = wb[tab["name"]]
        if tab.get("columns") == []:
            build_dashboard(ws, tab, wb)
        else:
            build_data_tab(ws, tab)
            add_charts(ws, tab, wb)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
