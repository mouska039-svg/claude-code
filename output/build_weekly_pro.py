#!/usr/bin/env python3
"""Builder PRO pour le planner hebdomadaire de dépenses — 52 semaines + rollup mensuel."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

C = {
    "hdr_dark":  "1C3A2F",
    "hdr_mid":   "2D5C48",
    "sage":      "5C7A6B",
    "sage_lt":   "A8C4B8",
    "terra":     "D4876B",
    "terra_lt":  "F0C9B7",
    "cream":     "F5EFE6",
    "cream2":    "EDE5D8",
    "white":     "FFFFFF",
    "ink":       "2C2C2A",
    "mist":      "8B8A87",
    "input_bg":  "FFFDF7",
    "sun":       "FFF3E0",   # dimanche highlight
    "sat":       "F3F8F5",   # samedi highlight
}

def fill(h): return PatternFill("solid", fgColor=h)
def font(c="2C2C2A", sz=10, bold=False, italic=False, name="Calibri"):
    return Font(color=c, size=sz, bold=bold, italic=italic, name=name)
def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def border_thin(color="D4C9B8"):
    s = Side(style="thin", color=color)
    return Border(top=s, bottom=s, left=s, right=s)
def border_outer():
    s = Side(style="medium", color="1C3A2F")
    return Border(top=s, bottom=s, left=s, right=s)

CATS = [
    ("🏠 Logement",     207),   # 900/month ÷ 4.33
    ("🛒 Alimentation", 104),
    ("🚗 Transport",     46),
    ("💊 Santé",         23),
    ("🎮 Loisirs",       35),
    ("📱 Abonnements",   14),
    ("👗 Shopping",      28),
    ("💡 Divers",        18),
]
DAYS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
DAYS_SHORT = ["Lun", "Mar", "Mer", "Jeu", "Ven", "Sam", "Dim"]

# Cols: A(pad) B(cat) C..I(days) J(total) K(budget) L(%) M(bar) N(status) O(pad)
# Indices: B=2 C=3 D=4 E=5 F=6 G=7 H=8 I=9  J=10 K=11 L=12 M=13 N=14
WCOLS = [3, 20, 11, 11, 11, 11, 11, 11, 11, 15, 15, 9, 24, 8, 3]

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def rept(row):
    return (f"=IF(K{row}=0,\"—\","
            f"REPT(\"▓\",MIN(20,INT(J{row}/K{row}*20)))&"
            f"REPT(\"░\",MAX(0,20-MIN(20,INT(J{row}/K{row}*20)))))")

def status(row):
    return (f"=IF(J{row}=0,\"⬜\","
            f"IF(J{row}/K{row}<0.8,\"🟢\","
            f"IF(J{row}/K{row}<=1,\"🟡\",\"🔴\")))")

def build_week(ws, week_num, year=2026):
    """Build one weekly tab — S01…S52."""
    set_widths(ws, WCOLS)
    ws.row_dimensions[1].height = 8

    # ── Title ──
    ws.merge_cells("B2:N2")
    c = ws["B2"]
    c.value = f"📅  SEMAINE {week_num:02d}  —  Suivi de dépenses hebdomadaire"
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 13, True)
    c.alignment = align("center"); ws.row_dimensions[2].height = 34

    # ── Sub-header: days ──
    ws.merge_cells("B3:B3")  # cat label
    c = ws["B3"]; c.value = "Catégorie"
    c.fill = fill(C["sage"]); c.font = font(C["white"], 9, True)
    c.alignment = align("center"); c.border = border_thin()

    day_bgs = {6: C["sat_lt"] if "sat_lt" in C else C["cream2"],
               7: C["sun"]}
    for di, day in enumerate(DAYS_SHORT):
        col = 3 + di
        bg = C["sun"] if di == 6 else (C["sat"] if di == 5 else C["sage"])
        fg = C["terra"] if di == 6 else (C["sage_lt"] if di == 5 else C["white"])
        c = ws.cell(row=3, column=col, value=day)
        c.fill = fill(bg); c.font = font(fg, 9, True)
        c.alignment = align("center"); c.border = border_thin()

    for col, label in [(10,"Total sem."),(11,"Budget sem."),(12,"% utilisé"),
                       (13,"▓▓▓ Progression"),(14,"🚦")]:
        fg_c = C["white"] if col < 13 else C["white"]
        c = ws.cell(row=3, column=col, value=label)
        c.fill = fill(C["hdr_mid"]); c.font = font(fg_c, 9, True)
        c.alignment = align("center"); c.border = border_thin()
    ws.row_dimensions[3].height = 20

    # ── Data rows 4–11 ──
    for i, (cat, bud_week) in enumerate(CATS):
        r = 4 + i
        alt = C["cream2"] if i % 2 else C["white"]

        c = ws.cell(row=r, column=2, value=cat)
        c.fill = fill(alt); c.font = font(C["ink"], 10)
        c.alignment = align("left"); c.border = border_thin()

        for di in range(7):
            col = 3 + di
            bg = "FFF8F0" if di == 6 else (C["sat"] if di == 5 else C["input_bg"])
            c = ws.cell(row=r, column=col, value=0)
            c.fill = fill(bg)
            c.font = font(C["terra"] if di == 6 else C["ink"], 10)
            c.alignment = align("right"); c.border = border_thin()
            c.number_format = "#,##0 €"

        # Total = sum of days
        c = ws.cell(row=r, column=10, value=f"=SUM(C{r}:I{r})")
        c.fill = fill(alt); c.font = font(C["ink"], 10, True)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

        # Budget semaine
        c = ws.cell(row=r, column=11, value=bud_week)
        c.fill = fill(C["input_bg"]); c.font = font(C["ink"], 10)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

        # %
        c = ws.cell(row=r, column=12,
                    value=f"=IF(K{r}=0,0,J{r}/K{r})")
        c.fill = fill(alt); c.font = font(C["ink"], 10)
        c.alignment = align("center"); c.border = border_thin()
        c.number_format = "0%"

        # Barre
        c = ws.cell(row=r, column=13, value=rept(r))
        c.fill = fill(alt)
        c.font = Font(name="Courier New", size=10, color=C["sage"])
        c.alignment = align("left"); c.border = border_thin()

        # Status
        c = ws.cell(row=r, column=14, value=status(r))
        c.fill = fill(alt); c.font = font(C["ink"], 12)
        c.alignment = align("center"); c.border = border_thin()

        ws.row_dimensions[r].height = 20

    # ── Totals row 12 ──
    r = 12
    ws.row_dimensions[r].height = 22
    c = ws.cell(row=r, column=2, value="TOTAL SEMAINE")
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 10, True)
    c.alignment = align("left"); c.border = border_thin()

    for di in range(7):
        col = 3 + di
        c = ws.cell(row=r, column=col,
                    value=f"=SUM({get_column_letter(col)}4:{get_column_letter(col)}11)")
        c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 10, True)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

    c = ws.cell(row=r, column=10, value="=SUM(J4:J11)")
    c.fill = fill(C["terra"]); c.font = font(C["white"], 12, True)
    c.alignment = align("right"); c.border = border_thin()
    c.number_format = "#,##0 €"

    c = ws.cell(row=r, column=11, value="=SUM(K4:K11)")
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 10, True)
    c.alignment = align("right"); c.border = border_thin()
    c.number_format = "#,##0 €"

    c = ws.cell(row=r, column=12,
                value="=IF(K12=0,0,J12/K12)")
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 10, True)
    c.alignment = align("center"); c.border = border_thin()
    c.number_format = "0%"

    c = ws.cell(row=r, column=13, value=rept(r))
    c.fill = fill(C["hdr_dark"])
    c.font = Font(name="Courier New", size=10, color=C["sage_lt"], bold=True)
    c.alignment = align("left"); c.border = border_thin()

    c = ws.cell(row=r, column=14, value=status(r))
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 12)
    c.alignment = align("center"); c.border = border_thin()

    # Conditional: col J red if over budget
    ws.conditional_formatting.add(
        "J4:J11",
        FormulaRule(formula=["J4>K4"],
                    fill=fill("FDE8E0"),
                    font=Font(color="C0392B", bold=True, name="Calibri")))

    ws.freeze_panes = "C4"
    ws.sheet_view.showGridLines = False


# ── Monthly rollup (12 tabs) ─────────────────────────────────────────────────

# Which weeks belong to which month (approx. 2026 ISO calendar)
MONTH_WEEKS = {
    "Janvier":    list(range(1, 5)),
    "Fevrier":    list(range(5, 9)),
    "Mars":       list(range(9, 14)),
    "Avril":      list(range(14, 18)),
    "Mai":        list(range(18, 22)),
    "Juin":       list(range(22, 27)),
    "Juillet":    list(range(27, 31)),
    "Aout":       list(range(31, 35)),
    "Septembre":  list(range(35, 40)),
    "Octobre":    list(range(40, 44)),
    "Novembre":   list(range(44, 48)),
    "Decembre":   list(range(48, 53)),
}
MONTHS_LABEL = {k: v for k, v in zip(MONTH_WEEKS, [
    "Janvier","Février","Mars","Avril","Mai","Juin",
    "Juillet","Août","Septembre","Octobre","Novembre","Décembre"])}

def month_col_widths():
    return [3, 22, 16, 16, 16, 12, 24, 8, 3]

def build_month_rollup(ws, month_tab, weeks):
    set_widths(ws, month_col_widths())
    label = MONTHS_LABEL.get(month_tab, month_tab)
    ws.row_dimensions[1].height = 8

    ws.merge_cells("B2:H2")
    c = ws["B2"]
    c.value = f"📅  {label.upper()}  —  Récapitulatif mensuel (semaines {weeks[0]:02d}–{weeks[-1]:02d})"
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 12, True)
    c.alignment = align("center"); ws.row_dimensions[2].height = 32

    for col, h in zip(range(2, 9),
                      ["Catégorie","Total dépensé","Budget mensuel",
                       "Restant","% utilisé","▓▓▓ Progression","🚦"]):
        c = ws.cell(row=3, column=col, value=h)
        c.fill = fill(C["sage"]); c.font = font(C["white"], 9, True)
        c.alignment = align("center"); c.border = border_thin()
    ws.row_dimensions[3].height = 20

    for i, (cat, _) in enumerate(CATS):
        r = 4 + i
        alt = C["cream2"] if i % 2 else C["white"]

        c = ws.cell(row=r, column=2, value=cat)
        c.fill = fill(alt); c.font = font(C["ink"], 10)
        c.alignment = align("left"); c.border = border_thin()

        # Total dépensé = sum of J col (row 4+i) across relevant weeks
        week_sum = "+".join(f"S{w:02d}!J{4+i}" for w in weeks)
        c = ws.cell(row=r, column=3, value=f"={week_sum}")
        c.fill = fill(C["input_bg"]); c.font = font(C["ink"], 10)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

        # Budget mensuel = sum of K col (row 4+i) across weeks
        bud_sum = "+".join(f"S{w:02d}!K{4+i}" for w in weeks)
        c = ws.cell(row=r, column=4, value=f"={bud_sum}")
        c.fill = fill(C["cream2"] if i%2 else C["white"])
        c.font = font(C["ink"], 10)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

        # Restant
        c = ws.cell(row=r, column=5, value=f"=D{r}-C{r}")
        c.fill = fill(alt); c.font = font(C["ink"], 10)
        c.alignment = align("right"); c.border = border_thin()
        c.number_format = "#,##0 €"

        # %
        c = ws.cell(row=r, column=6, value=f"=IF(D{r}=0,0,C{r}/D{r})")
        c.fill = fill(alt); c.font = font(C["ink"], 10)
        c.alignment = align("center"); c.border = border_thin()
        c.number_format = "0%"

        # Barre (using C=dépensé D=budget)
        barre = (f"=IF(D{r}=0,\"—\","
                 f"REPT(\"▓\",MIN(20,INT(C{r}/D{r}*20)))&"
                 f"REPT(\"░\",MAX(0,20-MIN(20,INT(C{r}/D{r}*20)))))")
        c = ws.cell(row=r, column=7, value=barre)
        c.fill = fill(alt)
        c.font = Font(name="Courier New", size=10, color=C["sage"])
        c.alignment = align("left"); c.border = border_thin()

        # Status
        st = (f"=IF(C{r}=0,\"⬜\","
              f"IF(C{r}/D{r}<0.8,\"🟢\","
              f"IF(C{r}/D{r}<=1,\"🟡\",\"🔴\")))")
        c = ws.cell(row=r, column=8, value=st)
        c.fill = fill(alt); c.font = font(C["ink"], 12)
        c.alignment = align("center"); c.border = border_thin()

        ws.row_dimensions[r].height = 20

    # Totals r=12
    r = 12
    ws.row_dimensions[r].height = 22
    for col, val, fmt in [
        (2, "TOTAL MENSUEL", None),
        (3, "=SUM(C4:C11)", "#,##0 €"),
        (4, "=SUM(D4:D11)", "#,##0 €"),
        (5, "=SUM(E4:E11)", "#,##0 €"),
        (6, "=IF(D12=0,0,C12/D12)", "0%"),
    ]:
        c = ws.cell(row=r, column=col,
                    value=val if fmt != "0%" else val)
        c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 10, True)
        c.alignment = align("right" if col > 2 else "left")
        c.border = border_thin()
        if fmt: c.number_format = fmt

    barre12 = (f"=IF(D12=0,\"—\","
               f"REPT(\"▓\",MIN(20,INT(C12/D12*20)))&"
               f"REPT(\"░\",MAX(0,20-MIN(20,INT(C12/D12*20)))))")
    c = ws.cell(row=r, column=7, value=barre12)
    c.fill = fill(C["hdr_dark"])
    c.font = Font(name="Courier New", size=10, color=C["sage_lt"], bold=True)
    c.alignment = align("left"); c.border = border_thin()

    st12 = (f"=IF(C12=0,\"⬜\","
            f"IF(C12/D12<0.8,\"🟢\","
            f"IF(C12/D12<=1,\"🟡\",\"🔴\")))")
    c = ws.cell(row=r, column=8, value=st12)
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 12)
    c.alignment = align("center"); c.border = border_thin()

    ws.conditional_formatting.add(
        "E4:E11",
        FormulaRule(formula=["E4<0"],
                    fill=fill("FDE8E0"),
                    font=Font(color="C0392B", bold=True, name="Calibri")))

    ws.freeze_panes = "B4"
    ws.sheet_view.showGridLines = False


# ── Dashboard ─────────────────────────────────────────────────────────────────

def build_weekly_dashboard(ws):
    set_widths(ws, [3, 22, 18, 4, 22, 18, 4, 22, 18, 3])
    ws.row_dimensions[1].height = 8

    ws.merge_cells("B2:I2")
    c = ws["B2"]
    c.value = "💰  WEEKLY EXPENSE PLANNER  —  Vue d'ensemble annuelle"
    c.fill = fill(C["hdr_dark"]); c.font = font(C["white"], 14, True)
    c.alignment = align("center"); ws.row_dimensions[2].height = 40

    ws.merge_cells("B3:I3")
    c = ws["B3"]
    c.value = "Toutes les données ci-dessous sont automatiques"
    c.fill = fill(C["hdr_mid"]); c.font = font(C["sage_lt"], 9, italic=True)
    c.alignment = align("center"); ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 10

    ws.merge_cells("B5:I5")
    c = ws["B5"]
    c.value = "  📊  INDICATEURS CLÉS ANNUELS"
    c.fill = fill(C["sage"]); c.font = font(C["white"], 10, True)
    c.alignment = align("left"); ws.row_dimensions[5].height = 24
    ws.row_dimensions[6].height = 8

    month_tabs = list(MONTH_WEEKS.keys())

    def kpi(ws, tl_col, tl_row, label, value, fmt, accent):
        br_row = tl_row + 1
        ws.merge_cells(f"{get_column_letter(tl_col)}{tl_row}:{get_column_letter(tl_col+1)}{tl_row}")
        c = ws.cell(row=tl_row, column=tl_col, value=label)
        c.fill = fill(accent); c.font = font(C["white"], 9)
        c.alignment = align("center"); ws.row_dimensions[tl_row].height = 22

        ws.merge_cells(f"{get_column_letter(tl_col)}{br_row}:{get_column_letter(tl_col+1)}{br_row}")
        c = ws.cell(row=br_row, column=tl_col, value=value)
        c.fill = fill(accent)
        c.font = Font(name="Calibri", size=20, bold=True, color=C["white"])
        c.alignment = align("center"); c.number_format = fmt
        ws.row_dimensions[br_row].height = 32

    total_dep = "+".join(f"{mt}!C12" for mt in month_tabs)
    total_bud = "+".join(f"{mt}!D12" for mt in month_tabs)

    kpi(ws, 2, 7, "💸  Total dépensé annuel",
        f"={total_dep}", "#,##0 €", C["terra"])
    kpi(ws, 5, 7, "📅  Budget annuel total",
        f"={total_bud}", "#,##0 €", C["hdr_mid"])
    kpi(ws, 8, 7, "💚  Économies estimées",
        f"={total_bud}-({total_dep})", "#,##0 €", "27AE60")

    ws.row_dimensions[9].height = 10

    kpi(ws, 2, 10, "📉  % Budget consommé",
        f"=IFERROR(({total_dep})/({total_bud}),0)", "0.0%", C["sage"])
    kpi(ws, 5, 10, "📆  Semaines saisies",
        "=COUNTA(S01:S52!J12)-COUNTIF(S01:S52!J12,0)", "0", C["hdr_dark"])
    kpi(ws, 8, 10, "🔢  Moy. / semaine",
        f"=IFERROR(({total_dep})/52,0)", "#,##0 €", C["terra"])

    ws.row_dimensions[12].height = 14

    # Charts section
    ws.merge_cells("B13:I13")
    c = ws["B13"]
    c.value = "  📈  DÉPENSES PAR MOIS  (automatique)"
    c.fill = fill(C["sage"]); c.font = font(C["white"], 10, True)
    c.alignment = align("left"); ws.row_dimensions[13].height = 24
    ws.row_dimensions[14].height = 8

    # Hidden data for chart
    months_short = ["Jan","Fév","Mar","Avr","Mai","Jun",
                    "Jul","Aoû","Sep","Oct","Nov","Déc"]
    ws.cell(row=30, column=2, value="Mois")
    ws.cell(row=30, column=3, value="Dépensé")
    ws.cell(row=30, column=4, value="Budget")
    for i, (mt, ms) in enumerate(zip(month_tabs, months_short)):
        r = 31 + i
        ws.cell(row=r, column=2, value=ms)
        ws.cell(row=r, column=3, value=f"='{mt}'!C12")
        ws.cell(row=r, column=4, value=f"='{mt}'!D12")
        ws.row_dimensions[r].height = 0
    ws.row_dimensions[30].height = 0

    bar = BarChart()
    bar.type = "col"; bar.grouping = "clustered"
    bar.title = "Dépensé vs Budget par mois"
    bar.style = 10; bar.width = 28; bar.height = 14
    data = Reference(ws, min_col=3, max_col=4, min_row=30, max_row=42)
    cats = Reference(ws, min_col=2, min_row=31, max_row=42)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.series[0].graphicalProperties.solidFill = C["terra"]
    bar.series[1].graphicalProperties.solidFill = C["hdr_mid"]
    ws.add_chart(bar, "B15")

    ws.sheet_view.showGridLines = False


# ── Main ──────────────────────────────────────────────────────────────────────

def build(out_path):
    wb = Workbook()
    wb.remove(wb.active)

    # Dashboard first for easy access
    ws_dash = wb.create_sheet("Dashboard")

    # 12 monthly rollup tabs
    month_sheets = {}
    for mt in MONTH_WEEKS:
        ws = wb.create_sheet(mt)
        month_sheets[mt] = ws

    # 52 weekly tabs
    week_sheets = {}
    for w in range(1, 53):
        ws = wb.create_sheet(f"S{w:02d}")
        week_sheets[w] = ws

    # Build
    for w in range(1, 53):
        build_week(week_sheets[w], w)

    for mt, weeks in MONTH_WEEKS.items():
        build_month_rollup(month_sheets[mt], mt, weeks)

    build_weekly_dashboard(ws_dash)

    # Tab colors
    ws_dash.sheet_properties.tabColor = C["terra"]
    for ws in month_sheets.values():
        ws.sheet_properties.tabColor = C["sage"]
    for ws in week_sheets.values():
        ws.sheet_properties.tabColor = C["hdr_mid"]

    wb.save(out_path)
    print(out_path)


if __name__ == "__main__":
    import sys
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/test_weekly.xlsx")
